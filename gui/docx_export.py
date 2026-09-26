import os
from pathlib import Path
from docx import Document
from docx.shared import Pt, Inches, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

def set_cell_border(cell, **kwargs):
    """
    Set cell borders for docx table
    kwargs: top, bottom, left, right
    values: dict(sz=12, val='single', color='AAAAAA')
    """
    tcPr = cell._tc.get_or_add_tcPr()
    tcBorders = tcPr.first_child_found_in("w:tcBorders")
    if tcBorders is None:
        tcBorders = OxmlElement('w:tcBorders')
        tcPr.append(tcBorders)
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        edge_data = kwargs.get(edge)
        if edge_data:
            tag = 'w:{}'.format(edge)
            element = tcBorders.find(qn(tag))
            if element is None:
                element = OxmlElement(tag)
                tcBorders.append(element)
            for key, attr in [("val", "w:val"), ("color", "w:color"), ("sz", "w:sz"), ("space", "w:space")]:
                if key in edge_data:
                    element.set(qn(attr), str(edge_data[key]))

class DocxExporter:
    def __init__(self, font_name="游明朝", font_size=10.5):
        self.font_name = font_name
        self.font_size = font_size

    def export(self, pages_data, output_path, options=None):
        """
        pages_data: list of dict:
        [
            {
                "page_num": int,
                "headings": list of str,
                "body_text": str,
                "tables": list of list of list of str (table[row][col]),
                "annotations": list of str,
                "figures": list of str (image paths)
            },
            ...
        ]
        """
        options = options or {}
        insert_page_break = options.get("insert_page_break", True)
        include_headings = options.get("include_headings", True)
        include_tables = options.get("include_tables", True)
        include_annotations = options.get("include_annotations", True)
        include_figures = options.get("include_figures", True)
        annotation_mode = options.get("annotation_numbering", "通し番号")

        doc = Document()
        
        # Set normal style font
        style = doc.styles['Normal']
        font = style.font
        font.name = self.font_name
        font.size = Pt(self.font_size)
        font.color.rgb = RGBColor(0x33, 0x33, 0x33)
        # For Asian fonts in docx
        rPr = style.element.get_or_add_rPr()
        rFonts = OxmlElement('w:rFonts')
        rFonts.set(qn('w:eastAsia'), self.font_name)
        rFonts.set(qn('w:ascii'), self.font_name)
        rFonts.set(qn('w:hAnsi'), self.font_name)
        rPr.append(rFonts)

        # Set page margins
        sections = doc.sections
        for section in sections:
            section.top_margin = Inches(0.8)
            section.bottom_margin = Inches(0.8)
            section.left_margin = Inches(0.9)
            section.right_margin = Inches(0.9)

        total_pages = len(pages_data)
        global_anno_idx = 1

        for idx, page in enumerate(pages_data):
            page_num = page.get("page_num", idx + 1)
            
            # Headings
            if include_headings and page.get("headings"):
                for h in page["headings"]:
                    if h.strip():
                        p = doc.add_paragraph()
                        p.paragraph_format.space_before = Pt(8)
                        p.paragraph_format.space_after = Pt(4)
                        p.paragraph_format.keep_with_next = True
                        run = p.add_run(h.strip())
                        run.bold = True
                        run.font.size = Pt(self.font_size + 2.5)
                        run.font.name = self.font_name
                        run.font.color.rgb = RGBColor(0x11, 0x11, 0x33)

            # Body Text
            body = page.get("body_text", "")
            if body:
                paragraphs = body.split("\n")
                for p_text in paragraphs:
                    if p_text.strip():
                        p = doc.add_paragraph()
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(3)
                        p.paragraph_format.line_spacing = 1.15
                        run = p.add_run(p_text.strip())
                        run.font.size = Pt(self.font_size)
                        run.font.name = self.font_name

            # Tables
            if include_tables and page.get("tables"):
                for tbl_data in page["tables"]:
                    if not tbl_data:
                        continue
                    rows = len(tbl_data)
                    cols = max(len(r) for r in tbl_data) if rows > 0 else 0
                    if rows == 0 or cols == 0:
                        continue
                    
                    doc_table = doc.add_table(rows=rows, cols=cols)
                    doc_table.style = 'Table Grid'
                    doc_table.autofit = True
                    
                    for r_idx, row in enumerate(tbl_data):
                        for c_idx, val in enumerate(row):
                            if c_idx < cols:
                                cell = doc_table.cell(r_idx, c_idx)
                                cell.text = str(val).strip()
                                for p in cell.paragraphs:
                                    p.paragraph_format.space_before = Pt(1)
                                    p.paragraph_format.space_after = Pt(1)
                                    p.paragraph_format.line_spacing = 1.0
                                    for run in p.runs:
                                        run.font.size = Pt(self.font_size - 1.0)
                                        run.font.name = self.font_name
                                        if r_idx == 0:
                                            run.bold = True
                    doc.add_paragraph() # space after table

            # Figures
            if include_figures and page.get("figures"):
                for fig_path in page["figures"]:
                    if os.path.exists(fig_path):
                        try:
                            p = doc.add_paragraph()
                            p.alignment = WD_ALIGN_PARAGRAPH.CENTER
                            run = p.add_run()
                            run.add_picture(fig_path, width=Inches(4.5))
                            p.paragraph_format.space_before = Pt(4)
                            p.paragraph_format.space_after = Pt(4)
                        except Exception as e:
                            print(f"[DocxExporter] Figure embed failed: {e}")

            # Annotations
            if include_annotations and page.get("annotations"):
                annos = [a.strip() for a in page["annotations"] if a.strip()]
                if annos:
                    p = doc.add_paragraph()
                    p.paragraph_format.space_before = Pt(8)
                    p.paragraph_format.space_after = Pt(2)
                    run = p.add_run("【注釈】")
                    run.bold = True
                    run.font.size = Pt(self.font_size - 1.0)
                    
                    for a_idx, anno_text in enumerate(annos):
                        p = doc.add_paragraph()
                        p.paragraph_format.space_before = Pt(0)
                        p.paragraph_format.space_after = Pt(1)
                        p.paragraph_format.left_indent = Inches(0.2)
                        
                        prefix = f"*{global_anno_idx} " if annotation_mode == "通し番号" else f"*{a_idx+1} "
                        run = p.add_run(f"{prefix}{anno_text}")
                        run.font.size = Pt(self.font_size - 1.5)
                        run.font.name = self.font_name
                        run.font.color.rgb = RGBColor(0x55, 0x55, 0x55)
                        global_anno_idx += 1

            # Page Break
            if insert_page_break and idx < total_pages - 1:
                doc.add_page_break()

        output_path = Path(output_path)
        output_path.parent.mkdir(parents=True, exist_ok=True)
        doc.save(str(output_path))
        return str(output_path)
